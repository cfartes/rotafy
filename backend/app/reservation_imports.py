"""Prévia CSV/XLSX limitada, sem fórmulas, seguida por confirmação atômica."""

import csv
import io
import json
import zipfile
from datetime import timedelta
from uuid import UUID, uuid4, uuid5
from xml.etree.ElementTree import ParseError

from defusedxml.common import DefusedXmlException
from fastapi import APIRouter, HTTPException, Request, UploadFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError
from sqlalchemy import text
from starlette.concurrency import run_in_threadpool

from app.billing.calendar import utc_now
from app.billing.service import lock_company
from app.identity import Auth, audit, scope
from app.operations import detail as trip_detail
from app.operations import rows, tenant
from app.reservations import (
    Passenger,
    Reservation,
    bookable,
    create_reservation,
    expire_holds,
    fingerprint,
    validate_passengers,
)

router = APIRouter(prefix="/api/v1", tags=["Importação de reservas"])
MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 500
LEGACY_HEADERS = ["familia", "responsavel", "telefone", "pais", "passageiro", "cpf"]
HEADERS = LEGACY_HEADERS + ["estrangeiro", "menor_5_anos", "nascimento"]


def parse_file(content, filename):
    if not content or len(content) > MAX_BYTES:
        raise HTTPException(413, "Envie um arquivo não vazio de até 2 MB.")
    try:
        if filename.lower().endswith(".csv"):
            decoded = content.decode("utf-8-sig")
            sample = decoded.splitlines()[0]
            delimiter = ";" if sample.count(";") > sample.count(",") else ","
            reader = csv.reader(io.StringIO(decoded), delimiter=delimiter, strict=True)
            values = []
            for row in reader:
                if len(row) > len(HEADERS):
                    raise ValueError("Quantidade de colunas inválida.")
                values.append(row)
                if len(values) > MAX_ROWS + 1:
                    raise HTTPException(
                        413, "O limite é de 500 passageiros por arquivo."
                    )
        elif filename.lower().endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                files = archive.infolist()
                if (
                    len(files) > 500
                    or sum(f.file_size for f in files) > 20 * 1024 * 1024
                ):
                    raise HTTPException(
                        413, "Planilha excede o limite de conteúdo descompactado."
                    )
                if any("vbaProject" in f.filename or f.flag_bits & 1 for f in files):
                    raise ValueError(
                        "Macros e arquivos criptografados não são aceitos."
                    )
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            try:
                if len(workbook.worksheets) != 1:
                    raise ValueError("Use uma única aba.")
                sheet = workbook.worksheets[0]
                sheet.reset_dimensions()
                values = []
                for index, row in enumerate(sheet.iter_rows(), 1):
                    if index > MAX_ROWS + 1:
                        raise HTTPException(
                            413, "O limite é de 500 passageiros por arquivo."
                        )
                    if any(c.data_type == "f" for c in row):
                        raise ValueError(
                            "Fórmulas não são aceitas. Cole somente valores."
                        )
                    if any(c.value is not None for c in row[len(HEADERS) :]):
                        raise ValueError("Quantidade de colunas inválida.")
                    if any(
                        c.value is not None and not isinstance(c.value, str)
                        for c in row[: len(HEADERS)]
                    ):
                        raise ValueError(
                            "Use células de texto, inclusive para telefone e CPF."
                        )
                    values.append(
                        [c.value or "" for c in row[: len(HEADERS)]]
                        + [""] * max(0, len(HEADERS) - len(row))
                    )
            finally:
                workbook.close()
        else:
            raise HTTPException(422, "Use CSV UTF-8 ou XLSX, conforme o modelo.")
        header = [str(value).strip().lower() for value in values[0]] if values else []
        while header and not header[-1]:
            header.pop()
        if header not in [LEGACY_HEADERS, HEADERS]:
            raise ValueError("Cabecalho diferente do modelo.")
        records = []
        for row in values[1:]:
            if any(value for value in row[len(header) :]):
                raise ValueError("Quantidade de colunas invalida.")
            record = dict(
                zip(
                    header,
                    row[: len(header)] + [""] * max(0, len(header) - len(row)),
                    strict=True,
                )
            )
            records.append({key: record.get(key, "") for key in HEADERS})
        return records
    except HTTPException:
        raise
    except (
        UnicodeDecodeError,
        csv.Error,
        zipfile.BadZipFile,
        KeyError,
        ValueError,
        TypeError,
        IndexError,
        InvalidFileException,
        ParseError,
        DefusedXmlException,
        OSError,
    ):
        # Não retornar valores de células, exceções XML ou detalhes internos.
        raise HTTPException(
            422,
            "Arquivo inválido. Use o cabeçalho familia,responsavel,telefone,pais,passageiro,cpf; uma aba e células de texto, sem fórmulas.",
        ) from None


def import_flag(value):
    normalized = value.strip().lower()
    if normalized in ("", "nao", "n\u00e3o", "false", "0"):
        return False
    if normalized in ("sim", "true", "1"):
        return True
    raise ValueError("Use sim ou nao nas flags de estrangeiro/menor de 5 anos.")


def normalize_families(records):
    grouped = {}
    errors = []
    for line, row in enumerate(records, 2):
        if not any(value.strip() for value in row.values()):
            continue
        reference = row["familia"].strip()
        try:
            if not reference or len(reference) > 100:
                raise ValueError(
                    "Identificador de família obrigatório, até 100 caracteres."
                )
            record = Reservation(
                operation_key=uuid4(),
                contact_name=row["responsavel"],
                phone=row["telefone"],
                phone_country=row["pais"].strip().upper() or None,
                passengers=[
                    Passenger(
                        full_name=row["passageiro"],
                        cpf=row["cpf"],
                        is_foreign=import_flag(row.get("estrangeiro", "")),
                        under_five=import_flag(row.get("menor_5_anos", "")),
                        birth_date=row.get("nascimento") or None,
                    )
                ],
            )
            item = grouped.setdefault(
                reference,
                {
                    "reference": reference,
                    "contact_name": record.contact_name,
                    "phone": record.phone,
                    "passengers": [],
                    "lines": [],
                },
            )
            if (item["contact_name"].casefold(), item["phone"]) != (
                record.contact_name.casefold(),
                record.phone,
            ):
                raise ValueError(
                    "A mesma família deve ter o mesmo responsável e telefone em todas as linhas."
                )
            item["passengers"].append(record.passengers[0].model_dump(mode="json"))
            item["lines"].append(line)
        except ValidationError as error:
            # Pydantic.errors inclui os valores de entrada: devolver somente mensagens.
            errors.append(
                {
                    "line": line,
                    "message": "; ".join(
                        e["msg"].removeprefix("Value error, ") for e in error.errors()
                    ),
                }
            )
        except ValueError as error:
            errors.append({"line": line, "message": str(error)})
    families = []
    for item in grouped.values():
        try:
            Reservation(
                operation_key=uuid4(),
                contact_name=item["contact_name"],
                phone=item["phone"],
                passengers=item["passengers"],
            )
            signatures = [
                (p["full_name"].casefold(), p["cpf"]) for p in item["passengers"]
            ]
            if len(set(signatures)) != len(signatures):
                raise ValueError(
                    "Passageiro repetido na família. Corrija ou diferencie seu nome no arquivo."
                )
            families.append(item)
        except ValidationError as error:
            errors.append(
                {
                    "line": item["lines"][0],
                    "message": "; ".join(
                        e["msg"].removeprefix("Value error, ") for e in error.errors()
                    ),
                }
            )
        except ValueError as error:
            errors.append({"line": item["lines"][0], "message": str(error)})
    if not families and not errors:
        errors.append({"line": 2, "message": "O arquivo não contém passageiros."})
    return sorted(families, key=lambda f: f["reference"]), errors


async def import_detail(db, import_id):
    found = await rows(
        db, "SELECT * FROM reservation_imports WHERE id=:id", id=import_id
    )
    if not found:
        raise HTTPException(404, "Prévia não encontrada nesta empresa.")
    result = found[0]
    result["reservations"] = await rows(
        db,
        "SELECT id,contact_name,passenger_count,status,import_reference FROM reservations WHERE import_id=:id ORDER BY import_reference",
        id=import_id,
    )
    return result


@router.post("/trips/{trip_id}/reservation-imports/preview", status_code=201)
async def preview(trip_id: UUID, request: Request, actor: Auth, file: UploadFile):
    tenant(actor, write=True)
    try:
        content = await file.read(MAX_BYTES + 1)
        records = await run_in_threadpool(parse_file, content, file.filename or "")
        families, errors = await run_in_threadpool(normalize_families, records)
    finally:
        await file.close()
    signature = fingerprint(records)
    async with request.app.state.platform_engine.begin() as db:
        await scope(db, actor.company_id)
        await lock_company(db, actor.company_id, require_active=True)
        trip = await trip_detail(db, trip_id)
        bookable(trip)
        valid_families = []
        for family in families:
            try:
                validate_passengers(trip, family["passengers"])
                valid_families.append(family)
            except HTTPException as error:
                errors.append({"line": family["lines"][0], "message": error.detail})
        families = valid_families
        prior = await db.scalar(
            text(
                "SELECT id FROM reservation_imports WHERE trip_id=:trip AND fingerprint=:hash"
            ),
            {"trip": trip_id, "hash": signature},
        )
        if prior:
            await db.execute(
                text(
                    "UPDATE reservation_imports SET expires_at=:expires WHERE id=:id AND status='preview' AND expires_at<=:now"
                ),
                {
                    "id": prior,
                    "now": utc_now(),
                    "expires": utc_now() + timedelta(minutes=30),
                },
            )
            return await import_detail(db, prior)
        import_id = uuid4()
        await db.execute(
            text("""INSERT INTO reservation_imports(id,company_id,trip_id,fingerprint,families,errors,expires_at)
          VALUES(:id,:company,:trip,:hash,CAST(:families AS jsonb),CAST(:errors AS jsonb),:expires)"""),
            {
                "id": import_id,
                "company": actor.company_id,
                "trip": trip_id,
                "hash": signature,
                "families": json.dumps(families),
                "errors": json.dumps(errors),
                "expires": utc_now() + timedelta(minutes=30),
            },
        )
        await audit(
            db,
            actor.company_id,
            actor.user["id"],
            "reservation_import.previewed",
            import_id,
        )
        return await import_detail(db, import_id)


@router.post("/reservation-imports/{import_id}/apply")
async def apply_import(import_id: UUID, request: Request, actor: Auth):
    tenant(actor, write=True)
    async with request.app.state.platform_engine.begin() as db:
        await scope(db, actor.company_id)
        await lock_company(db, actor.company_id, require_active=True)
        await expire_holds(db, actor.company_id)
        preview = await import_detail(db, import_id)
        if preview["status"] == "applied":
            return preview
        if preview["errors"]:
            raise HTTPException(
                409, "Corrija todos os erros no arquivo antes de importar."
            )
        if preview["expires_at"] <= utc_now():
            raise HTTPException(
                409,
                "A prévia expirou. Envie o arquivo novamente para atualizar a prévia.",
            )
        trip = await trip_detail(db, preview["trip_id"])
        bookable(trip)
        for family in preview["families"]:
            if await db.scalar(
                text(
                    "SELECT id FROM reservations WHERE trip_id=:trip AND import_reference=:reference"
                ),
                {"trip": trip["id"], "reference": family["reference"]},
            ):
                raise HTTPException(
                    409,
                    f"A família {family['reference']} já foi importada nesta viagem. Nenhuma reserva foi criada.",
                )
            data = Reservation(
                operation_key=uuid5(import_id, family["reference"]),
                contact_name=family["contact_name"],
                phone=family["phone"],
                passengers=family["passengers"],
            )
            try:
                await create_reservation(
                    db, actor, trip, data, import_id, family["reference"]
                )
            except HTTPException as error:
                raise HTTPException(
                    error.status_code,
                    f"Família {family['reference']}: {error.detail} Nenhuma reserva desta importação foi criada.",
                ) from None
        await db.execute(
            text(
                "UPDATE reservation_imports SET status='applied',applied_at=now() WHERE id=:id"
            ),
            {"id": import_id},
        )
        await audit(
            db,
            actor.company_id,
            actor.user["id"],
            "reservation_import.applied",
            import_id,
        )
        return await import_detail(db, import_id)
