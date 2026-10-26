"""Gera HTML portátil e apresentação editável a partir das fontes revisáveis."""

import base64
import hashlib
import html
import json
import re
from pathlib import Path

import markdown
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "publicacao"
OUTPUT = ROOT / "frontend" / "public" / "documentacao"
NAVY = "0B253F"
TEAL = "0C7A76"
MUTED = "577084"
DATE = "26/10/2026"


def image_url(name):
    path = ROOT / "docs" / "screenshots" / name
    return "data:image/png;base64," + base64.b64encode(path.read_bytes()).decode()


def manual(
    name="manual",
    title="Manual da empresa",
    audience="Administradores, atendentes, motoristas e guias",
):
    source = (SOURCE / f"{name}.md").read_text(encoding="utf-8")
    source = re.sub(
        r"\[CAPTURA:([\w.-]+)\]",
        lambda match: (
            '<figure><img src="'
            + image_url(match[1])
            + '" alt="Tela demonstrativa do Nexus Rotafy"/><figcaption>Tela demonstrativa; os menus dependem do perfil e da empresa.</figcaption></figure>'
        ),
        source,
    )
    engine = markdown.Markdown(
        extensions=["tables", "toc", "attr_list"],
        extension_configs={"toc": {"toc_depth": "2"}},
    )
    content = engine.convert(source)
    chapters = re.split(r"(?=<h2\b)", content)
    content = "".join(
        '<section class="chapter">' + part + "</section>"
        for part in chapters
        if part.strip()
    )
    template = (SOURCE / "manual-template.html").read_text(encoding="utf-8")
    template = template.replace("{{CONTENT}}", content).replace("{{TOC}}", engine.toc)
    template = (
        template.replace("{{NAME}}", name)
        .replace("{{TITLE}}", title)
        .replace("{{AUDIENCE}}", audience)
        .replace(
            "{{HELP}}", "/nexus/ajuda" if name == "manual-superadmin" else "/ajuda"
        )
        .replace("{{COUNT}}", str(len(re.findall(r"^## ", source, re.MULTILINE))))
    )
    (OUTPUT / f"{name}.html").write_text(template, encoding="utf-8")


def presentation(data):
    cards = []
    count = len(data["slides"])
    for index, slide in enumerate(data["slides"], 1):
        screenshot = (
            '<figure><img src="'
            + image_url(slide["image"])
            + '" alt="Tela demonstrativa do sistema"/></figure>'
            if slide.get("image")
            else ""
        )
        cards.append(
            f'<section class="slide {"with-image" if screenshot else ""} {"hero" if index in {1, count} else ""}" id="slide-{index}" aria-label="Slide {index} de {count}">'
            + '<div class="slide-top"><span>rotafy<span class="brand-dot">.</span> <small>NEXUS</small></span>'
            + f"<span>{index:02d} / {count:02d}</span></div>"
            + '<div class="slide-grid"><div><p class="eyebrow">'
            + html.escape(slide["eyebrow"])
            + "</p><h1>"
            + html.escape(slide["title"]).replace("\n", "<br>")
            + '</h1><p class="lead">'
            + html.escape(slide["lead"])
            + "</p><ul>"
            + "".join(
                "<li>" + html.escape(point) + "</li>" for point in slide["points"]
            )
            + "</ul></div>"
            + screenshot
            + "</div><footer>"
            + html.escape(slide["footnote"])
            + "</footer></section>"
        )
    template = (SOURCE / "presentation-template.html").read_text(encoding="utf-8")
    (OUTPUT / "apresentacao.html").write_text(
        template.replace("{{SLIDES}}", "".join(cards)), encoding="utf-8"
    )


def text_box(slide, x, y, w, h, text, size, color, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_after = Pt(11)
    return box


def powerpoint(data):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.core_properties.title = "Nexus Rotafy — apresentação do sistema"
    deck.core_properties.subject = "Funcionalidades e funcionamento · edição 2026.10"
    deck.core_properties.author = "Nexus Data Analytics"
    for index, item in enumerate(data["slides"], 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        hero = index in {1, len(data["slides"])}
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor.from_string(NAVY if hero else "F4F7FA")
        color = "FFFFFF" if hero else NAVY
        stripe = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), deck.slide_height
        )
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = RGBColor.from_string(TEAL)
        stripe.line.fill.background()
        width = 8 if item.get("image") else 11.8
        text_box(slide, 0.65, 0.35, 3, 0.4, "rotafy.  /  NEXUS", 18, color, True)
        text_box(
            slide,
            11.4,
            0.4,
            1.3,
            0.35,
            f"{index:02d} / {len(data['slides']):02d}",
            12,
            color,
        )
        text_box(
            slide,
            0.65,
            1.0,
            width,
            0.4,
            item["eyebrow"],
            11,
            "54D6CB" if hero else TEAL,
            True,
        )
        text_box(
            slide,
            0.65,
            1.52,
            width,
            1.28,
            item["title"],
            31 if item.get("image") else 36,
            color,
            True,
        )
        text_box(
            slide, 0.65, 2.9, width, 0.72, item["lead"], 20, "C3E4EA" if hero else MUTED
        )
        text_box(
            slide,
            0.7,
            3.85,
            width,
            2.4,
            "\n".join("• " + point for point in item["points"]),
            18,
            color,
        )
        text_box(
            slide,
            0.65,
            6.8,
            12,
            0.45,
            item["footnote"],
            11,
            "C3E4EA" if hero else MUTED,
        )
        if item.get("image"):
            path = ROOT / "docs" / "screenshots" / item["image"]
            with Image.open(path) as picture:
                ratio = picture.width / picture.height
            width_image = min(3.5, 5.2 * ratio)
            height_image = width_image / ratio
            slide.shapes.add_picture(
                str(path),
                Inches(9.1 + (3.5 - width_image) / 2),
                Inches(1.1),
                width=Inches(width_image),
                height=Inches(height_image),
            )
        slide.notes_slide.notes_text_frame.text = item["notes"]
    private_output = ROOT / "backend" / "private" / "documentacao"
    private_output.mkdir(parents=True, exist_ok=True)
    deck.save(private_output / "apresentacao.pptx")
    (OUTPUT / "apresentacao.pptx").unlink(missing_ok=True)



def update_documentation_csp():
    """Allow only the exact bundled documentation scripts, not arbitrary inline JS."""
    hashes = set()
    for name in ("manual", "manual-superadmin", "apresentacao"):
        markup = (OUTPUT / f"{name}.html").read_text(encoding="utf-8")
        for script in re.findall(r"<script>(.*?)</script>", markup, re.DOTALL):
            value = base64.b64encode(hashlib.sha256(script.encode("utf-8")).digest()).decode()
            hashes.add("'sha256-" + value + "'")
    config = ROOT / "frontend" / "nginx.conf"
    content = config.read_text(encoding="utf-8")
    content = re.sub(r"script-src[^;]+;", "script-src 'self' " + " ".join(sorted(hashes)) + ";", content)
    config.write_text(content, encoding="utf-8")


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    data = json.loads((SOURCE / "apresentacao.json").read_text(encoding="utf-8"))
    manual()
    manual(
        "manual-superadmin",
        "Manual do superadmin Nexus",
        "Administração da plataforma: empresas, planos, assinaturas e comunicação",
    )
    presentation(data)
    powerpoint(data)
    update_documentation_csp()
    print(
        f"HTML do manual e da apresentação gerados; {len(data['slides'])} slides editáveis."
    )


if __name__ == "__main__":
    main()
